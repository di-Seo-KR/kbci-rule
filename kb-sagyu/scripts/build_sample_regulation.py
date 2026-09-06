# -*- coding: utf-8 -*-
"""사규관리시스템 개발계 신구조문대비표 기능 테스트용 샘플 사규 2종(현행/개정안) + 기대 대비표.
「사규관리규정」(2026. 8. 12.) 작성기준을 그대로 따른 가상 사규 「샘플규정」."""
import os, re, shutil, difflib, copy
import docx
from docx.shared import Pt, Mm, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = '/home/user/kbci-rule/kb-sagyu/sources/(11-2)_사규관리규정_260812.docx'   # 신 템플릿 서식 본
F_TITLE, F_MED, F_LIGHT = "KB금융 제목체 Medium", "KB금융 본문체 Medium", "KB금융 본문체 Light"
BLUE = RGBColor(0x00, 0x00, 0xCD); BLACK = RGBColor(0, 0, 0)
MARK = re.compile(r'\(\d{4}\. \d{1,2}\. \d{1,2}\. (?:개정|신설|삭제)\)|<\d{4}\. \d{1,2}\. \d{1,2}\. (?:개정|신설)>')


def kfont(run, size=11, bold=False, color=BLACK, font=F_LIGHT):
    run.font.size = Pt(size); run.font.bold = bold; run.font.name = font; run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr(); rf = rPr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rPr.append(rf)
    for a in ('w:eastAsia', 'w:ascii', 'w:hAnsi'): rf.set(qn(a), font)


def runs(p, text, size=11, bold=False, font=F_LIGHT):
    """개정 표기 (일자 개정/신설/삭제)·<일자 개정> 만 파란색."""
    pos = 0
    for m in MARK.finditer(text):
        if m.start() > pos: kfont(p.add_run(text[pos:m.start()]), size, bold, BLACK, font)
        kfont(p.add_run(m.group(0)), size, False, BLUE, font); pos = m.end()
    if pos < len(text): kfont(p.add_run(text[pos:]), size, bold, BLACK, font)


# ─────────── 조문 데이터: (종류, 텍스트)  종류 ch=장, a=조, bu=부칙표기, bua=부칙조, tbl=별표
CUR = [
 ("ch", "제1장 총칙"),
 ("a", "제1조(목적) 이 규정은 사규관리시스템의 신구조문대비표 기능을 검증하기 위한 샘플 사규로서, 사규의 작성 형식과 개정 표기 방식을 시험하는 것을 목적으로 한다."),
 ("a", "제2조(정의) 이 규정에서 사용하는 용어의 정의는 다음과 같다.\n1. “샘플”이란 실제 업무에 적용되지 아니하는 시험용 문서를 말한다.\n2. “테스트”란 사규관리시스템의 기능을 확인하는 행위를 말한다.\n3. “검증자”란 테스트를 수행하는 사규관리부서의 담당자를 말한다."),
 ("a", "제3조(적용범위) 이 규정은 사규관리시스템 개발계에서 수행하는 테스트에 한하여 적용한다."),
 ("ch", "제2장 테스트의 수행"),
 ("a", "제4조(테스트 계획) ① 검증자는 테스트를 시작하기 전에 테스트 계획을 수립하여야 한다.\n② 테스트 계획에는 대상 기능, 수행 일정 및 판정 기준을 포함한다."),
 ("a", "제5조(테스트 수행) ① 검증자는 테스트 계획에 따라 기능을 확인하고 그 결과를 기록한다.\n② 결과 기록에는 다음 각 호의 사항을 포함한다.\n1. 확인한 기능의 명칭\n2. 확인 일시 및 확인자\n3. 정상 여부 및 오류 내용"),
 ("a", "제6조(오류의 처리) 검증자는 오류를 발견한 경우 다음 각 호의 구분에 따라 처리한다.\n1. 경미한 오류 : 개발자에게 통보\n2. 중대한 오류 : 테스트 중단 후 사규관리부서장에게 보고\n3. 판단이 어려운 오류 : 검증자 간 협의"),
 ("a", "제7조(결과 보고) 검증자는 테스트를 마친 때에는 [서식 제1호]의 테스트 결과 보고서를 작성하여 사규관리부서장에게 보고하여야 한다."),
 ("ch", "제3장 보칙"),
 ("a", "제8조(판정 기준) 기능별 판정 기준은 [별표 제1호]와 같다."),
 ("a", "제9조(기록의 보존) 테스트 결과 기록은 1년간 보존한다."),
 ("a", "제10조(재검토기한) 회사는 이 규정에 대하여 2026년 1월 1일을 기준으로 매 3년이 되는 시점마다 그 타당성을 검토하여 개선 등의 조치를 하여야 한다."),
 ("bu", "부 칙(2026. 1. 1. 제정)"),
 ("bua", "제1조(시행일) 이 규정은 2026. 1. 1. 부터 시행한다."),
]
CUR_TABLE = ("[별표 제1호] 기능별 판정 기준", [["구분", "판정 기준", "비고"],
                                             ["조문 비교", "변경 문장이 파란색으로 표시될 것", ""],
                                             ["신설 조문", "현행란에 “신설”로 표시될 것", ""],
                                             ["삭제 조문", "개정란에 “삭제”로 표시될 것", ""]])
CUR_FORM = "[서식 제1호] 테스트 결과 보고서"
CUR_HEAD = "[시행 2026. 1. 1.] [2026. 1. 1. 제정-2026-001]"

NEW = [
 ("ch", "제1장 총칙"),
 ("a", "제1조(목적) 이 규정은 사규관리시스템의 신구조문대비표 기능을 검증하기 위한 샘플 사규로서, 사규의 작성 형식과 개정 표기 방식을 시험하는 것을 목적으로 한다."),
 ("a", "제2조(정의) 이 규정에서 사용하는 용어의 정의는 다음과 같다.\n1. “샘플”이란 실제 업무에 적용되지 아니하는 시험용 문서를 말한다.\n2. “테스트”란 사규관리시스템의 기능을 확인하는 행위를 말한다.\n3. “검증자”란 테스트를 수행하는 사규관리부서의 담당자와 사규관리부서장이 지정하는 자를 말한다. (2026. 9. 15. 개정)\n4. “개발계”란 운영계와 분리되어 기능 개발 및 시험에 사용하는 사규관리시스템 환경을 말한다. (2026. 9. 15. 신설)"),
 ("a", "제3조(적용범위) 이 규정은 사규관리시스템 개발계에서 수행하는 테스트에 한하여 적용한다."),
 ("ch", "제2장 테스트의 수행"),
 ("a", "제4조(테스트 계획의 수립) ① 검증자는 테스트를 시작하기 전에 테스트 계획을 수립하여 사규관리부서장의 승인을 받아야 한다. (2026. 9. 15. 개정)\n② 테스트 계획에는 대상 기능, 수행 일정 및 판정 기준을 포함한다."),
 ("a", "제5조(테스트 수행) ① 검증자는 테스트 계획에 따라 기능을 확인하고 그 결과를 기록한다.\n② 결과 기록에는 다음 각 호의 사항을 포함한다.\n1. 확인한 기능의 명칭\n2. 확인 일시 및 확인자\n3. 정상 여부 및 오류 내용\n③ 검증자는 결과 기록을 사규관리시스템에 등록하여야 한다. (2026. 9. 15. 신설)"),
 ("a", "제6조(오류의 처리) 검증자는 오류를 발견한 경우 다음 각 호의 구분에 따라 처리한다.\n1. 경미한 오류 : 개발자에게 통보\n2. (2026. 9. 15. 삭제)\n3. 판단이 어려운 오류 : 검증자 간 협의"),
 ("a", "제7조(결과 보고) 검증자는 테스트를 마친 때에는 [서식 제1호]의 테스트 결과 보고서를 작성하여 사규관리부서장에게 보고하여야 한다."),
 ("a", "제7조의2(재테스트) 사규관리부서장은 보고받은 오류가 조치된 경우 검증자에게 재테스트를 지시할 수 있다. (2026. 9. 15. 신설)"),
 ("ch", "제3장 보칙"),
 ("a", "제8조(판정 기준) 기능별 판정 기준은 [별표 제1호]와 같다."),
 ("a", "제9조 (2026. 9. 15. 삭제)"),
 ("a", "제10조(재검토기한) 회사는 이 규정에 대하여 2026년 1월 1일을 기준으로 매 3년이 되는 시점마다 그 타당성을 검토하여 개선 등의 조치를 하여야 한다."),
 ("bu", "부 칙(2026. 1. 1. 제정)"),
 ("bua", "제1조(시행일) 이 규정은 2026. 1. 1. 부터 시행한다."),
 ("bu", "부 칙(2026. 9. 15. 개정)"),
 ("bua", "제1조(시행일) 이 규정은 2026. 9. 15. 부터 시행한다."),
 ("bua", "제2조(경과조치) 이 규정 시행 전에 수립한 테스트 계획은 제4조제1항의 개정규정에 따라 승인을 받은 것으로 본다."),
]
NEW_TABLE = ("[별표 제1호] 기능별 판정 기준 <2026. 9. 15. 개정>", [["구분", "판정 기준", "비고"],
                                                            ["조문 비교", "변경 문장이 파란색으로 표시될 것", ""],
                                                            ["신설 조문", "현행란에 “신설”로 표시될 것", ""],
                                                            ["삭제 조문", "개정란에 “삭제”로 표시될 것", ""],
                                                            ["별표 비교", "별표 변경 시 제목 다음에 개정일자가 표시될 것", "2026. 9. 15. 추가"]])
NEW_FORM = "[서식 제1호] 테스트 결과 보고서"
NEW_HEAD = "[시행 2026. 9. 15.] [2026. 9. 15. 개정-2026-099]"


def build_doc(out, items, head, table, form):
    shutil.copy(BASE, out); doc = docx.Document(out); body = doc.element.body
    for ch in list(body.iterchildren()):
        if ch.tag != qn('w:sectPr'): body.remove(ch)
    hp = doc.sections[0].header.paragraphs[0]
    for r in list(hp.runs): r._element.getparent().remove(r._element)
    kfont(hp.add_run('샘플규정(소관부서: 경영전략부)'), size=9)

    def base_p(align=None, li=None, fi=None):
        p = doc.add_paragraph(); pf = p.paragraph_format
        pf.line_spacing = Pt(20); pf.space_before = Pt(0); pf.space_after = Pt(0)
        if align is not None: p.alignment = align
        if li is not None: pf.left_indent = Mm(li)
        if fi is not None: pf.first_line_indent = Mm(fi)
        return p

    p = base_p(AL.CENTER); kfont(p.add_run('샘플규정'), size=18, bold=True, font=F_TITLE)
    p = base_p(AL.CENTER); kfont(p.add_run(head), size=10, color=BLUE); base_p()
    for kind, text in items:
        if kind == 'ch':
            p = base_p(AL.CENTER); kfont(p.add_run(text), size=13, bold=True, font=F_MED); base_p()
        elif kind == 'bu':
            base_p(); p = base_p(AL.CENTER); kfont(p.add_run(text), size=12, bold=True, font=F_MED); base_p()
        elif kind == 'bua':
            runs(base_p(li=3.35, fi=-3.35), text)
        else:
            for i, ln in enumerate(text.split('\n')):
                if i == 0:
                    p = base_p(li=3.35, fi=-3.35); head_, _, rest = ln.partition(') ')
                    kfont(p.add_run(head_ + ')' if rest else ln), size=11, bold=True)
                    if rest: runs(p, ' ' + rest)
                elif ln[0] in '①②③④⑤⑥⑦⑧⑨⑩': runs(base_p(li=3.53), ln)
                else: runs(base_p(li=7.06), ln)
            base_p()
    # 별표
    base_p(); p = base_p(); runs(p, table[0], bold=False)
    t = doc.add_table(rows=len(table[1]), cols=len(table[1][0])); t.style = 'Table Grid'
    for ri, row in enumerate(table[1]):
        for ci, val in enumerate(row):
            c = t.rows[ri].cells[ci]; c.text = ''; pp = c.paragraphs[0]
            pp.alignment = AL.CENTER if (ri == 0 or ci != 1) else AL.LEFT
            kfont(pp.add_run(val), size=10, bold=(ri == 0))
    for ci, w in enumerate((Cm(3.2), Cm(9.0), Cm(3.5))):
        for r in t.rows: r.cells[ci].width = w
    base_p(); runs(base_p(), form); runs(base_p(li=3.53), '(서식 생략 — 테스트용)')
    doc.save(out); print('생성:', out)


# ─────────── 기대 신구조문대비표 (가로) ───────────
def arts(items):
    return {re.match(r'제\d+조(?:의\d+)?', t).group(0): t for k, t in items if k == 'a'}


def build_daebi(out):
    C, N = arts(CUR), arts(NEW)
    doc = docx.Document(); sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = Mm(297), Mm(210)
    sec.top_margin = sec.bottom_margin = Mm(14); sec.left_margin = sec.right_margin = Mm(13)
    ns = doc.styles['Normal']; ns.font.name = '맑은 고딕'; ns.font.size = Pt(9); ns.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    p = doc.add_paragraph(); p.alignment = 1; kfont(p.add_run('「샘플규정」 신구조문대비표 (기대 결과)'), size=15, bold=True, font='맑은 고딕')
    p = doc.add_paragraph(); p.alignment = 1; kfont(p.add_run('파란색 = 변경 문장 / 사규관리시스템 대비표 기능의 출력과 대조하는 용도'), size=8, color=RGBColor(0x77, 0x77, 0x77), font='맑은 고딕')
    t = doc.add_table(rows=1, cols=3); t.style = 'Table Grid'
    for j, h in enumerate(('현 행', '개 정 안', '비 고')):
        c = t.rows[0].cells[j]; c.text = ''; kfont(c.paragraphs[0].add_run(h), size=9.5, bold=True, font='맑은 고딕'); c.paragraphs[0].alignment = 1
    GREY = RGBColor(0x33, 0x33, 0x33)

    def render(cell, text, mask):
        cell.text = ''; first = True
        for ln in text.split('\n'):
            p = cell.paragraphs[0] if first else cell.add_paragraph(); first = False
            p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.15
            if ln[:1] in '①②③④⑤⑥⑦⑧⑨⑩': p.paragraph_format.left_indent = Mm(4)
            elif re.match(r'\d+\.', ln): p.paragraph_format.left_indent = Mm(8)
            base = text.index(ln) if ln else 0
            changed = mask is not None and any(mask[base:base + len(ln)])
            kfont(p.add_run(ln), size=9, color=BLUE if changed else GREY, font='맑은 고딕')

    keys = list(dict.fromkeys(list(C) + list(N)))
    keys.sort(key=lambda k: (int(re.search(r'\d+', k).group()), k))
    NOTE = {'제2조': '제3호 개정, 제4호 신설', '제4조': '제목 개정, 제1항 개정', '제5조': '제3항 신설',
            '제6조': '제2호 삭제', '제7조의2': '조 신설', '제9조': '조 삭제'}
    for k in keys:
        c, n = C.get(k), N.get(k)
        if c is None or n is None or c != n:
            row = t.add_row()
            for j, w in enumerate((Cm(11.5), Cm(11.5), Cm(4.0))): row.cells[j].width = w
            if c is None: render(row.cells[0], '〈신 설〉', None); render(row.cells[1], n, bytearray(b'\x01' * len(n)))
            else:
                mc, mn = bytearray(len(c)), bytearray(len(n))
                for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, c, n, autojunk=False).get_opcodes():
                    if tag in ('replace', 'delete'): mc[i1:i2] = b'\x01' * (i2 - i1)
                    if tag in ('replace', 'insert'): mn[j1:j2] = b'\x01' * (j2 - j1)
                render(row.cells[0], c, mc); render(row.cells[1], n, mn)
            row.cells[2].text = ''; kfont(row.cells[2].paragraphs[0].add_run(NOTE.get(k, '')), size=8, font='맑은 고딕')
    row = t.add_row()
    for j, w in enumerate((Cm(11.5), Cm(11.5), Cm(4.0))): row.cells[j].width = w
    render(row.cells[0], CUR_TABLE[0] + '\n(4행 표)', None); render(row.cells[1], NEW_TABLE[0] + '\n(5행 표 — “별표 비교” 행 추가)', bytearray(b'\x01' * 60))
    row.cells[2].text = ''; kfont(row.cells[2].paragraphs[0].add_run('별표 개정 — 제목 다음 <개정일자> 표기'), size=8, font='맑은 고딕')
    row = t.add_row()
    for j, w in enumerate((Cm(11.5), Cm(11.5), Cm(4.0))): row.cells[j].width = w
    render(row.cells[0], '〈신 설〉', None); render(row.cells[1], '부 칙(2026. 9. 15. 개정)\n제1조(시행일) 이 규정은 2026. 9. 15. 부터 시행한다.\n제2조(경과조치) …', bytearray(b'\x01' * 80))
    row.cells[2].text = ''; kfont(row.cells[2].paragraphs[0].add_run('부칙 신설'), size=8, font='맑은 고딕')
    doc.save(out); print('생성:', out, '| 행', len(t.rows) - 1)


build_doc('샘플규정_현행_260101.docx', CUR, CUR_HEAD, CUR_TABLE, CUR_FORM)
build_doc('샘플규정_개정안_260915.docx', NEW, NEW_HEAD, NEW_TABLE, NEW_FORM)
build_daebi('샘플규정_신구조문대비표_기대결과.docx')
