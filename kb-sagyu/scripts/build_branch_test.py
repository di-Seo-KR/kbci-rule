# -*- coding: utf-8 -*-
"""사규관리시스템 계층(가지번호) 처리 시험용 샘플 사규 생성.

    python3 kb-sagyu/scripts/build_branch_test.py [출력폴더]

산출: ① 정상 표기본  ② 오표기본(시스템이 잡는지 확인용)  ③ 개정안(가지번호 신설·삭제)
실제 사규에 존재하는 패턴만 사용 — 조 가지번호 제1조의2~제31조의10, 호 가지번호 1의2·2의2·10의2,
오표기 제25조2·제19조2·제20조의1·14-2. 등.
"""
import os, re, sys, shutil
import docx
from docx.shared import Pt, Mm, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TPL = os.path.join(BASE, 'sources/(11-2)_사규관리규정_260812.docx')
F_TITLE, F_MED, F_LIGHT = "KB금융 제목체 Medium", "KB금융 본문체 Medium", "KB금융 본문체 Light"
BLUE = RGBColor(0x00, 0x00, 0xCD); BLACK = RGBColor(0, 0, 0)
MARK = re.compile(r'\(\d{4}\. \d{1,2}\. \d{1,2}\. (?:개정|신설|삭제)\)|<\d{4}\. \d{1,2}\. \d{1,2}\. (?:개정|신설)>')
HANG = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮"
MOK = "가나다라마바사아자차카타파하"


def kfont(run, size=11, bold=False, color=BLACK, font=F_LIGHT):
    run.font.size = Pt(size); run.font.bold = bold; run.font.name = font; run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr(); rf = rPr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rPr.append(rf)
    for a in ('w:eastAsia', 'w:ascii', 'w:hAnsi'): rf.set(qn(a), font)


def runs(p, text, size=11, bold=False, font=F_LIGHT):
    pos = 0
    for m in MARK.finditer(text):
        if m.start() > pos: kfont(p.add_run(text[pos:m.start()]), size, bold, BLACK, font)
        kfont(p.add_run(m.group(0)), size, False, BLUE, font); pos = m.end()
    if pos < len(text): kfont(p.add_run(text[pos:]), size, bold, BLACK, font)


def level(line):
    """0=조 두문, 1=항, 2=호, 3=목"""
    s = line.strip()
    if s[:1] in HANG: return 1
    if re.match(r'^\d+(?:의\d+)?\.', s) or re.match(r'^\d+-\d+\.', s): return 2
    if re.match(r'^[' + MOK + r'](?:의\d+)?\.', s): return 3
    return 0


def build(items, head, title, out, tables=(), forms=()):
    shutil.copy(TPL, out); doc = docx.Document(out); body = doc.element.body
    for ch in list(body.iterchildren()):
        if ch.tag != qn('w:sectPr'): body.remove(ch)
    hp = doc.sections[0].header.paragraphs[0]
    for r in list(hp.runs): r._element.getparent().remove(r._element)
    kfont(hp.add_run(title + '(소관부서: 경영전략부)'), size=9)

    def P(align=None, li=None, fi=None):
        p = doc.add_paragraph(); pf = p.paragraph_format
        pf.line_spacing = Pt(20); pf.space_before = Pt(0); pf.space_after = Pt(0)
        if align is not None: p.alignment = align
        if li is not None: pf.left_indent = Mm(li)
        if fi is not None: pf.first_line_indent = Mm(fi)
        return p

    p = P(AL.CENTER); kfont(p.add_run(title), size=18, bold=True, font=F_TITLE)
    p = P(AL.CENTER); kfont(p.add_run(head), size=10, color=BLUE); P()
    for kind, text in items:
        if kind == 'ch':
            p = P(AL.CENTER); kfont(p.add_run(text), size=13, bold=True, font=F_MED); P()
        elif kind == 'se':
            p = P(AL.CENTER); kfont(p.add_run(text), size=12, bold=True, font=F_MED); P()
        elif kind == 'bu':
            P(); p = P(AL.CENTER); kfont(p.add_run(text), size=12, bold=True, font=F_MED); P()
        elif kind == 'bua':
            runs(P(li=3.35, fi=-3.35), text)
        else:
            for i, ln in enumerate(text.split('\n')):
                if i == 0:
                    p = P(li=3.35, fi=-3.35); h, _, rest = ln.partition(') ')
                    kfont(p.add_run(h + ')' if rest else ln), size=11, bold=True)
                    if rest: runs(p, ' ' + rest)
                else:
                    runs(P(li=(0, 3.53, 7.06, 10.59)[level(ln)]), ln.strip())
            P()
    if tables or forms:
        P()
        for t in tables: runs(P(), t)
        for f in forms: runs(P(), f)
    doc.save(out); print('생성:', os.path.basename(out))


# ══════════════════ ① 정상 표기본 ══════════════════
OK_ITEMS = [
("ch", "제1장 총칙"),
("a", "제1조(목적) 이 규정은 사규관리시스템이 장·절·관·조·항·호·목과 가지번호를 올바르게 인식하고 정렬하는지 시험하기 위한 표본으로서, 「사규관리규정」 제13조부터 제15조까지의 표기 방식을 모두 담는 것을 목적으로 한다."),
("a", """제2조(정의) 이 규정에서 사용하는 용어의 정의는 다음과 같다.
1. “호”란 조 또는 항에 여러 사항을 적을 때 쓰는 항목을 말한다.
1의2. “호의 가지번호”란 기존 호 사이에 새 호를 넣을 때 쓰는 번호를 말한다.
1의3. “연속 가지번호”란 가지번호가 둘 이상 이어지는 것을 말한다.
2. “목”이란 하나의 호에 여러 사항을 적을 때 쓰는 항목을 말하며, 다음 각 목과 같이 표기한다.
가. 첫째 목
나. 둘째 목
나의2. 목의 가지번호
다. 셋째 목
2의2. “목의 가지번호”란 기존 목 사이에 새 목을 넣을 때 쓰는 번호를 말한다.
3. “정렬”이란 조·호·목을 번호 순서대로 늘어놓는 것을 말한다.
10. “열 번째 호”란 두 자리 호 번호를 시험하기 위한 항목을 말한다.
10의2. “두 자리 호의 가지번호”란 두 자리 호에 붙는 가지번호를 말한다."""),
("a", "제3조(적용범위) 이 규정은 사규관리시스템 개발계의 계층 인식 시험에만 적용한다."),
("a", "제3조의2(다른 사규와의 관계) 사규의 표기에 관하여는 「사규관리규정」에서 정하는 바에 따른다."),

("ch", "제2장 항·호·목의 단계"),
("a", "제4조(항의 표기) " + "\n".join(
    [HANG[0] + " 항은 ①부터 순서대로 붙이며, 제2항부터 줄을 바꾸어 적는다."] +
    ["%s 제%d항이다. 두 자리 항 번호가 바르게 표시되는지 확인한다." % (HANG[i], i + 1) for i in range(1, 15)])),
("a", """제5조(호의 표기) 호는 다음 각 호와 같이 적는다.
1. 첫째 호
2. 둘째 호
3. 셋째 호
4. 넷째 호
5. 다섯째 호
6. 여섯째 호
7. 일곱째 호
8. 여덟째 호
9. 아홉째 호
10. 열째 호
11. 열한째 호
12. 열두째 호"""),
("a", "제6조(목의 표기) 목은 다음 각 목과 같이 적는다.\n1. 목의 시험\n" + "\n".join(
    "%s. %s째 목" % (MOK[i], "첫둘셋넷다섯여섯일곱여덟아홉열열한열두열셋열넷"[i]) for i in range(14))),
("a", """제7조(네 단계의 중첩) ① 조·항·호·목이 네 단계로 겹치는 경우를 시험한다.
② 다음 각 호와 같다.
1. 첫째 호
가. 첫째 목
나. 둘째 목
2. 둘째 호
가. 첫째 목
나. 둘째 목
다. 셋째 목"""),

("ch", "제2장의2 조 가지번호"),
("se", "제1절 연속 가지번호"),
("a", "제7조의2(가지번호의 시작) 조 사이에 새 조를 넣을 때에는 앞 조의 번호에 가지번호를 붙인다."),
("a", "제7조의3(연속) 가지번호가 이어질 때에는 제7조의2 다음에 제7조의3을 붙인다."),
("a", "제7조의4(정렬) 가지번호가 붙은 조는 제7조와 제8조 사이에 놓인다."),
("a", "제7조의5(인용) 가지번호가 붙은 조를 인용할 때에는 제7조의2와 같이 적는다."),
("a", "제7조의6(삭제 대상) 개정안에서 삭제되는 조로, 제정본에는 정상 조문으로 있다."),
("se", "제1절의2 두 자리 가지번호"),
("a", "제7조의7(두 자리의 시작) 가지번호가 아홉을 넘으면 두 자리가 된다."),
("a", "제7조의8(중간) 제7조의7과 제7조의9 사이에 놓인다."),
("a", "제7조의9(한 자리의 끝) 한 자리 가지번호의 마지막이다."),
("a", "제7조의10(두 자리 ①) 제7조의10은 제7조의9 다음에 오고 제7조의2 앞에 오지 아니한다. 글자 순서로 늘어놓으면 제7조의10이 제7조의2보다 앞에 오므로, 번호 순서로 늘어놓는지 확인한다."),
("a", "제7조의11(두 자리 ②) 제7조의11은 이 장의 마지막 조이며, 다음 조는 제8조이다."),
("a", "제8조(가지번호 없는 조의 복귀) 가지번호가 붙은 조가 끝나면 다음 조는 제8조가 되며 제7조의12가 되지 아니한다."),

("ch", "제3장 인용"),
("a", """제9조(인용의 표기) ① 이 규정 안의 조문을 인용할 때에는 제7조의10제1항과 같이 적는다.
② 항·호·목을 함께 인용할 때에는 제7조제2항제1호가목과 같이 적는다.
③ 다른 사규를 인용할 때에는 「사규관리규정」 제14조와 같이 적고, 법령을 인용할 때에는 『상법』 제408조의2와 같이 적는다.
④ 별표와 서식은 [별표 제1호] 및 [서식 제1호]와 같이 적는다.
⑤ 제2조제1호의2 및 같은 조 제10호의2와 같이 호의 가지번호도 인용한다."""),

("ch", "제4장 보칙"),
("a", "제10조(위임) 이 규정의 시행에 필요한 세부사항은 사규관리부서장이 따로 정한다."),
("a", "제11조(재검토기한) 회사는 이 규정에 대하여 2026년 9월 15일을 기준으로 매 3년이 되는 시점마다 그 타당성을 검토하여 개선 등의 조치를 하여야 한다."),
("bu", "부 칙(2026. 9. 15. 제정)"),
("bua", "제1조(시행일) 이 규정은 2026. 9. 15. 부터 시행한다."),
]
OK_TABLES = ["[별표 제1호] 계층 표기 대조표", "[별표 제1호의2] 가지번호 대조표"]
OK_FORMS = ["[서식 제1호] 계층 인식 시험 결과서", "[서식 제1호의2] 정렬 시험 결과서"]

# ══════════════════ ② 오표기본 ══════════════════
NG_ITEMS = [
("ch", "제1장 총칙"),
("a", "제1조 (목적) 이 규정은 사규관리시스템이 잘못된 가지번호 표기를 찾아내는지 시험하기 위한 표본이다. 조와 괄호 사이에 공백이 있다."),
("a", """제2조 (정의) 이 규정에서 사용하는 용어의 정의는 다음과 같다.
1. “정상”이란 「사규관리규정」에 맞는 표기를 말한다.
1의2. “가지번호”란 정상 표기이다.
14-2. “하이픈 표기”란 호의 가지번호를 하이픈으로 적은 것으로 잘못된 표기이다.
17-2. “하이픈 표기 둘째”란 같은 잘못이 이어진 것을 말한다."""),
("a", "제3조2(의 누락) 조 가지번호에서 “의”가 빠진 잘못된 표기이다. 바른 표기는 제3조의2이다."),
("a", "제 7 조의 2(공백) 조·번호·가지번호 사이에 공백이 들어간 잘못된 표기이다. 바른 표기는 제7조의2이다."),
("a", "제7조의 3(가지번호 앞 공백) 가지번호 앞에 공백이 들어간 잘못된 표기이다. 바른 표기는 제7조의3이다."),
("a", "제19조2(의 누락 둘째) 실제 사규에서 발견된 잘못이다. 바른 표기는 제19조의2이다."),
("a", "제20조의1(가지번호 1) 가지번호는 2부터 시작하므로 제20조의1은 성립하지 아니한다."),
("a", "제25조2(의 누락 셋째) 「정관」에서 발견된 잘못이다. 바른 표기는 제25조의2이다."),
("bu", "부 칙 (2026. 09. 15.)"),
("bua", "① (시행일) 이 규정은 2026. 09. 15 부터 시행한다."),
]

# ══════════════════ ③ 개정안 (가지번호 신설·삭제) ══════════════════
def rev_items():
    it = []
    for k, t in OK_ITEMS:
        if k == 'a' and t.startswith('제7조의6('):
            it.append((k, "제7조의6 (2026. 10. 1. 삭제)"))
            it.append(("a", "제7조의6의2(삭제된 조 뒤의 신설) 삭제된 조 뒤에 가지번호를 더 붙인 경우를 시험한다. (2026. 10. 1. 신설)"))
            continue
        if k == 'a' and t.startswith('제3조의2('):
            it.append((k, "제3조의2(다른 사규와의 관계) 사규의 표기에 관하여는 「사규관리규정」에서 정하는 바에 따른다. (2026. 10. 1. 개정)"))
            it.append(("a", "제3조의3(신설된 연속 가지번호) 기존 제3조의2 다음에 새 가지번호를 붙인 경우를 시험한다. (2026. 10. 1. 신설)"))
            continue
        if k == 'a' and t.startswith('제2조(정의)'):
            it.append((k, t.replace('1의3. “연속 가지번호”란 가지번호가 둘 이상 이어지는 것을 말한다.',
                                    '1의3. “연속 가지번호”란 가지번호가 둘 이상 이어지는 것을 말한다.\n1의4. “신설된 호의 가지번호”란 개정으로 새로 넣은 호의 가지번호를 말한다. (2026. 10. 1. 신설)')
                        .replace('10의2. “두 자리 호의 가지번호”란 두 자리 호에 붙는 가지번호를 말한다.',
                                 '10의2. (2026. 10. 1. 삭제)')))
            continue
        if k == 'a' and t.startswith('제7조의11('):
            it.append((k, t))
            it.append(("a", "제7조의12(두 자리 가지번호의 신설) 제7조의11 다음에 새 가지번호를 붙인 경우를 시험한다. (2026. 10. 1. 신설)"))
            continue
        if k == 'bu':
            it.append(("bu", "부 칙(2026. 9. 15. 제정)")); it.append(("bua", "제1조(시행일) 이 규정은 2026. 9. 15. 부터 시행한다."))
            it.append(("bu", "부 칙(2026. 10. 1. 개정)")); it.append(("bua", "제1조(시행일) 이 규정은 2026. 10. 1. 부터 시행한다."))
            break
        it.append((k, t))
    return it


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, 'drafts/샘플')
    os.makedirs(out, exist_ok=True)
    build(OK_ITEMS, "[시행 2026. 9. 15.] [2026. 9. 15. 제정-2026-101]", "가지번호 시험 규정",
          os.path.join(out, "가지번호 시험 규정_정상표기_260915.docx"), OK_TABLES, OK_FORMS)
    build(NG_ITEMS, "[시행 2026. 09. 15.][2026. 09. 15. 제정-2026-102]", "가지번호 오표기 시험 규정",
          os.path.join(out, "가지번호 시험 규정_오표기_260915.docx"))
    build(rev_items(), "[시행 2026. 10. 1.] [2026. 10. 1. 개정-2026-103]", "가지번호 시험 규정",
          os.path.join(out, "가지번호 시험 규정_개정안_261001.docx"), OK_TABLES, OK_FORMS)
    return 0


if __name__ == '__main__':
    sys.exit(main())
