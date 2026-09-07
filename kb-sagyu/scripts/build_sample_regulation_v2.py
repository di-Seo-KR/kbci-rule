# -*- coding: utf-8 -*-
"""수정 업로드 후 신구대비표 기준본 버그 재현용 샘플 2종.
  ① 샘플규정_개정전문_수정업로드용_260915.docx : 2026. 9. 15. 개정본(개정-2026-099)을 '수정 기능'으로 다시 올릴 소폭 수정본
  ② 샘플규정_개정전문_v2_261001.docx           : ①을 기준으로 한 다음 개정안(2026. 10. 1. 개정-2026-100)
  + 기대 대비표 2종(정상: ①기준 / 버그: 수정 전 기준)"""
import re, os, sys
src = open('/home/user/kbci-rule/kb-sagyu/scripts/build_sample_regulation.py', encoding='utf-8').read()
exec(src.split("build_doc('샘플규정_현행")[0])      # 함수·데이터 정의만 가져온다
ORIG_NEW, ORIG_NEW_TABLE = list(NEW), NEW_TABLE   # 수정 전 260915본(전역 NEW는 대비표 생성 시 덮어씀)

def sub(items, key, old, new):
    out = []
    for k, t in items:
        if k == 'a' and t.startswith(key + '('):
            assert old in t, (key, old); t = t.replace(old, new)
        out.append((k, t))
    return out

# ── ① 수정업로드본: 개정 표기는 그대로, 문구만 소폭 손질(오탈자·표현 정정 성격)
FIX = list(NEW)
FIX = sub(FIX, '제1조', '시험하는 것을 목적으로 한다.', '검증하는 것을 목적으로 한다.')                       # X1 (v2에서 안 건드림)
FIX = sub(FIX, '제3조', '개발계에서 수행하는 테스트에 한하여 적용한다.', '개발계에서 수행하는 테스트 및 운영계 반영 전 검증에 적용한다.')  # Z (v2에서 또 개정)
FIX = sub(FIX, '제5조', '결과 기록을 사규관리시스템에 등록하여야 한다.', '결과 기록을 지체 없이 사규관리시스템에 등록하여야 한다.')      # X2 (v2에서 안 건드림)
FIX_TABLE = (NEW_TABLE[0], [r[:] for r in NEW_TABLE[1]]); FIX_TABLE[1][4][1] = '별표 변경 시 제목 다음에 <개정일자 개정>이 표시될 것'   # X3
FIX_HEAD = NEW_HEAD   # 같은 개정 건의 수정이므로 [시행 2026. 9. 15.] [2026. 9. 15. 개정-2026-099] 유지

# ── ② v2: 수정업로드본을 기준으로 한 다음 개정
V2 = list(FIX)
V2 = sub(V2, '제2조', '2. “테스트”란 사규관리시스템의 기능을 확인하는 행위를 말한다.',
         '2. “테스트”란 사규관리시스템의 기능을 확인하고 그 결과를 기록하는 행위를 말한다. (2026. 10. 1. 개정)')            # Y (수정업로드에서 안 건드린 조문)
V2 = sub(V2, '제3조', '테스트 및 운영계 반영 전 검증에 적용한다.', '테스트 및 운영계 반영 전 검증에 적용하며, 운영계에는 적용하지 아니한다. (2026. 10. 1. 개정)')   # Z
V2 = sub(V2, '제7조의2', '재테스트를 지시할 수 있다. (2026. 9. 15. 신설)',
         '재테스트를 지시할 수 있다. (2026. 9. 15. 신설)\n② 재테스트의 결과는 제7조에 따라 보고한다. (2026. 10. 1. 신설)')
V2 = [(k, t) for k, t in V2]
V2 = sub(V2, '제7조의2', '제7조의2(재테스트) 사규관리부서장은', '제7조의2(재테스트) ① 사규관리부서장은')
V2 += [("bu", "부 칙(2026. 10. 1. 개정)"), ("bua", "제1조(시행일) 이 규정은 2026. 10. 1. 부터 시행한다.")]
V2_TABLE = FIX_TABLE     # 별표는 v2에서 변경 없음
V2_HEAD = "[시행 2026. 10. 1.] [2026. 10. 1. 개정-2026-100]"

build_doc('샘플규정_개정전문_수정업로드용_260915.docx', FIX, FIX_HEAD, FIX_TABLE, NEW_FORM)
build_doc('샘플규정_개정전문_v2_261001.docx', V2, V2_HEAD, V2_TABLE, NEW_FORM)

# ── 기대 대비표: 정상(수정업로드본 기준) / 버그(수정 전 260915본 기준)
def daebi(out, cur_items, new_items, cur_tbl, new_tbl, title, note, notes, tail):
    """build_daebi(v1 전용 꼬리행·비고)를 v2 상황에 맞게 후처리한다."""
    global CUR, NEW, CUR_TABLE, NEW_TABLE
    CUR, NEW, CUR_TABLE, NEW_TABLE = cur_items, new_items, cur_tbl, new_tbl
    import docx as _d, re as _re
    build_daebi(out)
    d = _d.Document(out); GREY = RGBColor(0x33, 0x33, 0x33)
    for p, txt, sz, bold, col in ((d.paragraphs[0], title, 15, True, BLACK), (d.paragraphs[1], note, 8, False, RGBColor(0x77, 0x77, 0x77))):
        for r in list(p.runs): r._element.getparent().remove(r._element)
        kfont(p.add_run(txt), size=sz, bold=bold, color=col, font='맑은 고딕')
    t = d.tables[0]
    for row in t.rows[-2:]: row._tr.getparent().remove(row._tr)          # v1 전용 별표·부칙 행 제거
    for row in t.rows[1:]:                                              # 비고 재작성
        key = _re.match(r'제\d+조(?:의\d+)?', row.cells[1].text.strip() or row.cells[0].text.strip())
        row.cells[2].text = ''; kfont(row.cells[2].paragraphs[0].add_run(notes.get(key.group(0) if key else '', '')), size=8, font='맑은 고딕')
    for cur, new, memo in tail:                                         # 상황별 꼬리행
        row = t.add_row()
        for j, w in enumerate((Cm(11.5), Cm(11.5), Cm(4.0))): row.cells[j].width = w
        for cell, txt, blue in ((row.cells[0], cur, False), (row.cells[1], new, True)):
            cell.text = ''; first = True
            for ln in txt.split('\n'):
                p = cell.paragraphs[0] if first else cell.add_paragraph(); first = False
                p.paragraph_format.space_after = Pt(1)
                kfont(p.add_run(ln), size=9, bold=(ln.startswith('〈')), color=(RGBColor(0xB0, 0x30, 0x30) if ln.startswith('〈') else (BLUE if blue else GREY)), font='맑은 고딕')
        row.cells[2].text = ''; kfont(row.cells[2].paragraphs[0].add_run(memo), size=8, font='맑은 고딕')
    d.save(out)

BU = '부 칙(2026. 10. 1. 개정)\n제1조(시행일) 이 규정은 2026. 10. 1. 부터 시행한다.'
daebi('샘플규정_v2_대비표_정상(수정업로드본 기준).docx', FIX, V2, FIX_TABLE, V2_TABLE,
      '「샘플규정」 v2 신구조문대비표 — 정상 기대 결과 (수정업로드본 기준)',
      '현행란 = 수정업로드본(2026. 9. 15. 수정 반영). 제1조·제5조③·별표는 v2에서 변경이 없으므로 대비표에 나오지 않아야 함',
      {'제2조': '제2호 개정', '제3조': '개정 (현행란에 수정업로드 문구 “및 운영계 반영 전 검증에” 가 있어야 정상)', '제7조의2': '제1항 번호 부여, 제2항 신설'},
      [('〈신 설〉', BU, '부칙 신설')])
daebi('샘플규정_v2_대비표_버그시(수정 전 본문 기준).docx', ORIG_NEW, V2, ORIG_NEW_TABLE, V2_TABLE,
      '「샘플규정」 v2 신구조문대비표 — 버그 재현 시 나올 결과 (수정 전 본문 기준)',
      '현행란 = 수정 전 2026. 9. 15.본. 이 결과가 나오면 시스템이 수정업로드본이 아닌 수정 전 본문을 기준으로 삼은 것',
      {'제1조': '★ v2 변경 없음 — 수정업로드 반영분(“시험하는”→“검증하는”)이 v2 개정으로 잘못 잡힘',
       '제2조': '제2호 개정 (정상)', '제3조': '★ 현행란이 수정 전 문구(“테스트에 한하여”) — 수정업로드본 기준이면 “및 운영계 반영 전 검증에”여야 함',
       '제5조': '★ v2 변경 없음 — 수정업로드 반영분(“지체 없이” 추가)이 v2 개정으로 잘못 잡힘', '제7조의2': '제1항 번호 부여, 제2항 신설 (정상)'},
      [('[별표 제1호] 기능별 판정 기준 <2026. 9. 15. 개정>\n5행 “…제목 다음에 개정일자가 표시될 것”', '[별표 제1호] 기능별 판정 기준 <2026. 9. 15. 개정>\n5행 “…제목 다음에 <개정일자 개정>이 표시될 것”', '★ v2 변경 없음 — 수정업로드 반영분이 별표 개정으로 잘못 잡힘'),
       ('〈신 설〉', BU, '부칙 신설 (정상)')])
